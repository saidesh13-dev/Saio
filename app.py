import os
import re
from io import BytesIO
from pathlib import Path

import fitz
import pytesseract
from flask import Flask, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from pytesseract import Output
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config["UPLOAD_FOLDER"] = "uploads"
app.config["OUTPUT_FOLDER"] = "outputs"
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024

os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)
os.makedirs(app.config["OUTPUT_FOLDER"], exist_ok=True)


def normalize_cell(value):
    return "" if value is None else str(value).strip()


def split_row_cells(words):
    if not words:
        return []

    words = sorted(words, key=lambda item: item["x0"])
    cells = []
    current = [words[0]]
    last_x = words[0]["x0"]

    for word in words[1:]:
        if word["x0"] - last_x > 18:
            cells.append(current)
            current = [word]
        else:
            current.append(word)
        last_x = word["x0"]
    cells.append(current)

    result = []
    for cell in cells:
        text = " ".join(word["text"] for word in sorted(cell, key=lambda item: item["x0"]))
        if text.strip():
            result.append(text.strip())
    return result


def extract_form_rows(page):
    rows = []
    text_dict = page.get_text("dict")
    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue

        line_items = []
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                if text:
                    line_items.append({
                        "text": text,
                        "x0": span["bbox"][0],
                        "x1": span["bbox"][2],
                        "y0": span["bbox"][1],
                        "y1": span["bbox"][3],
                    })

        if line_items:
            rows.append({
                "Page": page.number + 1,
                "Y": round(min(item["y0"] for item in line_items) / 4) * 4,
                "Cells": split_row_cells(line_items),
                "Text": " ".join(item["text"] for item in sorted(line_items, key=lambda x: x["x0"]))
            })
    return rows


def get_fitz_word_items(page):
    items = []
    for word in page.get_text("words"):
        if len(word) < 5:
            continue
        x0, y0, x1, y1, text = word[:5]
        cleaned = str(text).strip()
        if not cleaned:
            continue
        items.append({
            "text": cleaned,
            "x0": float(x0),
            "x1": float(x1),
            "y0": float(y0),
            "y1": float(y1),
        })
    return items


def compact_layout_rows(rows):
    if not rows:
        return rows

    compacted = []
    for row in sorted(rows, key=lambda item: (item["Page"], item["Y"])):
        if not compacted:
            compacted.append(row)
            continue

        prev = compacted[-1]
        same_page = row["Page"] == prev["Page"]
        close_y = abs(row["Y"] - prev["Y"]) <= 3
        same_shape = len(row.get("Cells", [])) == len(prev.get("Cells", []))

        if same_page and close_y and same_shape:
            merged_cells = []
            for left, right in zip(prev.get("Cells", []), row.get("Cells", [])):
                merged_cells.append((left or "") if left else (right or ""))
            if len(row.get("Cells", [])) > len(prev.get("Cells", [])):
                merged_cells = row.get("Cells", [])
            if row.get("Text", "") and prev.get("Text", ""):
                combined_text = prev.get("Text", "") + " | " + row.get("Text", "")
            else:
                combined_text = prev.get("Text", "") or row.get("Text", "")
            prev["Cells"] = merged_cells
            prev["Text"] = combined_text
            prev["Y"] = min(prev["Y"], row["Y"])
        else:
            compacted.append(row)
    return compacted


def looks_like_amount(value):
    text = (value or "").strip()
    if not text:
        return False
    cleaned = text.replace("R", "").replace("$", "").replace(" ", "")
    if not cleaned:
        return False
    if re.fullmatch(r"[-+]?\d{1,3}(?:[.,]\d{3})*(?:[.,]\d{2})?", cleaned):
        return True
    if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", cleaned):
        return True
    if re.fullmatch(r"[-+]?\d+[.,]\d{2}[A-Za-z]*", cleaned):
        return True
    if re.fullmatch(r"[-+]?\d+[.,]\d{2}\s*[A-Za-z]+", cleaned):
        return True
    return bool(re.search(r"\d", text) and any(ch in text for ch in [".", ",", "-", "+", "R", "$", "C", "D"]))


def looks_like_date(value):
    text = (value or "").strip()
    if not text:
        return False
    return bool(re.search(r"\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}", text) or re.search(r"\d{2,4}[/\-.]\d{1,2}[/\-.]\d{1,2}", text))


def is_statement_like_page(page):
    words = get_fitz_word_items(page)
    if len(words) < 40:
        return False

    text_values = [word["text"].strip() for word in words]
    lower_text = {value.lower() for value in text_values}
    money_like = sum(
        1 for value in text_values
        if any(ch.isdigit() for ch in value) and ("R" in value or "." in value or "," in value or "/" in value or "-" in value)
    )
    date_like = sum(1 for value in text_values if looks_like_date(value))
    statement_headers = len(lower_text.intersection({
        "date", "transaction", "description", "amount", "balance",
        "transaction date", "transaction description",
    }))

    if statement_headers >= 2 and money_like >= 4:
        return True
    if date_like >= 3 and money_like >= 8:
        return True
    return False


def is_statement_document(doc):
    all_words = []
    for page in doc:
        all_words.extend(get_fitz_word_items(page))

    if len(all_words) < 40:
        return False

    text_values = [word["text"].strip() for word in all_words]
    lower_text = {value.lower() for value in text_values}
    statement_headers = len(lower_text.intersection({
        "date", "transaction", "description", "amount", "balance",
        "transaction date", "transaction description",
    }))
    money_like = sum(
        1 for value in text_values
        if any(ch.isdigit() for ch in value) and ("R" in value or "." in value or "," in value or "-" in value)
    )
    date_like = sum(1 for value in text_values if looks_like_date(value))

    return (statement_headers >= 2 and money_like >= 4) or (date_like >= 3 and money_like >= 8)


def extract_statement_rows(page):
    words = get_fitz_word_items(page)
    if not words:
        return []

    line_heights = sorted(item["y1"] - item["y0"] for item in words if item["y1"] > item["y0"])
    median_line_height = line_heights[len(line_heights) // 2] if line_heights else 10
    row_gap = max(3, min(7, median_line_height * 0.6))

    rows = []
    current_row = []
    current_y = None

    for word in sorted(words, key=lambda item: (item["y0"], item["x0"])):
        if current_y is None:
            current_y = word["y0"]
        if abs(word["y0"] - current_y) <= row_gap:
            current_row.append(word)
        else:
            if current_row:
                rows.append(current_row)
            current_row = [word]
            current_y = word["y0"]
    if current_row:
        rows.append(current_row)

    statement_rows = []
    for row_words in rows:
        row_words = sorted(row_words, key=lambda item: item["x0"])
        if not row_words:
            continue

        text_values = [item["text"].strip() for item in row_words if item["text"].strip()]
        if not text_values:
            continue

        if any(v.lower() in {"date", "transaction", "description", "amount", "balance"} for v in text_values):
            continue

        date_tokens = []
        desc_tokens = []
        amount_tokens = []
        balance_tokens = []
        amount_start = page.rect.width * 0.56
        balance_start = page.rect.width * 0.72

        for item in row_words:
            text = item["text"].strip()
            if not text:
                continue

            if item["x0"] < page.rect.width * 0.18 or looks_like_date(text):
                date_tokens.append(text)
            elif item["x0"] >= balance_start:
                balance_tokens.append(text)
            elif item["x0"] >= amount_start:
                amount_tokens.append(text)
            else:
                desc_tokens.append(text)

        if not amount_tokens and not balance_tokens:
            monetary_items = [item for item in row_words if looks_like_amount(item["text"])]
            monetary_items.sort(key=lambda item: item["x0"])
            if len(monetary_items) >= 2:
                amount_tokens = [monetary_items[-2]["text"]]
                balance_tokens = [monetary_items[-1]["text"]]
            elif monetary_items:
                amount_tokens = [monetary_items[0]["text"]]

        date = " ".join(date_tokens).strip()
        desc = " ".join(desc_tokens).strip()
        amount = amount_tokens[-1].strip() if amount_tokens else ""
        balance = balance_tokens[-1].strip() if balance_tokens else ""

        if not date and row_words:
            date = row_words[0]["text"].strip()

        cells = [date.strip(), " ".join(desc).strip(), amount.strip(), balance.strip()]
        if any(cell for cell in cells) and not all(cell == "" for cell in cells):
            statement_rows.append({
                "Page": page.number + 1,
                "Y": min(item["y0"] for item in row_words),
                "Cells": cells,
                "Text": " | ".join(cells),
            })

    return statement_rows


def is_scanned_or_stamped_page(page):
    words = page.get_text("words")
    image_count = len(page.get_images(full=True))
    if image_count and len(words) < 50:
        return True
    if len(words) < 12 and image_count >= 1:
        return True
    if len(words) < 8 and page.rect.width > 0:
        return True
    return False


def ocr_page_as_rows(page, page_number):
    try:
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        data = pytesseract.image_to_data(image, config="--psm 6", output_type=Output.DICT)

        words = []
        for idx, text in enumerate(data.get("text", [])):
            clean = text.strip()
            if not clean:
                continue
            conf = data.get("conf", ["-1"])[idx]
            try:
                conf_val = int(float(conf))
            except (ValueError, TypeError):
                conf_val = -1
            if conf_val >= 0 and conf_val < 25:
                continue

            left = float(data["left"][idx])
            top = float(data["top"][idx])
            width = float(data["width"][idx])
            height = float(data["height"][idx])
            words.append({
                "text": clean,
                "x0": left,
                "x1": left + width,
                "y0": top,
                "y1": top + height,
            })

        if not words:
            return []

        words = sorted(words, key=lambda item: (item["y0"], item["x0"]))
        rows = []
        current_row = []
        current_y = None

        for word in words:
            if current_y is None:
                current_row = [word]
                current_y = word["y0"]
                continue

            if abs(word["y0"] - current_y) <= 18:
                current_row.append(word)
            else:
                rows.append({
                    "Page": page_number,
                    "Y": current_y,
                    "Cells": split_row_cells(current_row),
                    "Text": " ".join(item["text"] for item in sorted(current_row, key=lambda x: x["x0"]))
                })
                current_row = [word]
                current_y = word["y0"]

        if current_row:
            rows.append({
                "Page": page_number,
                "Y": current_y,
                "Cells": split_row_cells(current_row),
                "Text": " ".join(item["text"] for item in sorted(current_row, key=lambda x: x["x0"]))
            })

        return rows
    except Exception:
        return []


def extract_pdf_structure(pdf_file):
    layout_rows = []
    image_rows = []

    doc = fitz.open(pdf_file)
    statement_document = is_statement_document(doc)
    for page_number, page in enumerate(doc, start=1):
        words = get_fitz_word_items(page)
        if statement_document:
            rows = extract_statement_rows(page)
        elif is_scanned_or_stamped_page(page):
            rows = ocr_page_as_rows(page, page_number)
        else:
            rows = extract_form_rows(page)

        for row in rows:
            layout_rows.append({
                "Page": page_number,
                "Y": row["Y"],
                "Cells": row.get("Cells", [row.get("Text", "")]),
                "Text": row.get("Text", ""),
            })

        for image_index, image in enumerate(page.get_images(full=True), start=1):
            try:
                xref = image[0]
                pix = fitz.Pixmap(doc, xref)
                if pix.n == 4:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                image_rows.append({
                    "Page": page_number,
                    "ImageIndex": image_index,
                    "XRef": xref,
                    "Pixmap": pix,
                })
            except Exception:
                continue

    return compact_layout_rows(layout_rows), image_rows


def build_workbook_from_pdf(pdf_file):
    workbook = Workbook()
    form_sheet = workbook.active
    form_sheet.title = "Form_Data"

    layout_rows, image_rows = extract_pdf_structure(pdf_file)
    if not layout_rows and not image_rows:
        raise ValueError("No readable text or images were found in the PDF.")

    if layout_rows:
        max_columns = max((len(item["Cells"]) for item in layout_rows), default=1)
        if len(layout_rows[0].get("Cells", [])) == 4 and any(
            looks_like_date(str(cell)) for cell in layout_rows[0].get("Cells", [])
        ):
            header = ["Date", "Description", "", "", "", "Amount", "Balance"]
        else:
            header = ["Page", "Row", "Text"] + [f"Column_{i}" for i in range(1, max_columns + 1)]
        form_sheet.append(header)
        for row_index, item in enumerate(layout_rows, start=1):
            row_values = item["Cells"] + [""] * (max_columns - len(item["Cells"]))
            if header == ["Date", "Description", "", "", "", "Amount", "Balance"]:
                form_sheet.append([row_values[0], row_values[1], "", "", "", row_values[2], row_values[3]])
            else:
                form_sheet.append([item["Page"], row_index, item["Text"], *row_values[:max_columns]])

    layout_sheet = workbook.create_sheet("Layout")
    layout_sheet.append(["Page", "Row", "Text", "Y Position"])
    for row_index, item in enumerate(layout_rows, start=1):
        layout_sheet.append([item["Page"], row_index, item["Text"], item["Y"]])

    if image_rows:
        image_sheet = workbook.create_sheet("Images")
        image_sheet.append(["Page", "Image Index", "Image File"])

        image_dir = Path(app.config["UPLOAD_FOLDER"]) / "images"
        image_dir.mkdir(exist_ok=True, parents=True)

        for row_index, image_info in enumerate(image_rows, start=2):
            page_number = image_info["Page"]
            image_index = image_info["ImageIndex"]
            pix = image_info["Pixmap"]
            file_name = f"page_{page_number}_image_{image_index}.png"
            file_path = image_dir / file_name
            pix.save(file_path)

            image_sheet.cell(row=row_index, column=1, value=page_number)
            image_sheet.cell(row=row_index, column=2, value=image_index)
            image_sheet.cell(row=row_index, column=3, value=file_name)

            try:
                image_cell = XLImage(str(file_path))
                image_sheet.add_image(image_cell, f"F{row_index}")
            except Exception:
                pass

    return workbook


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        file = request.files.get("pdf_file")
        if not file or file.filename == "":
            return render_template("index.html", error="Please choose a PDF file to upload.")

        if not file.filename.lower().endswith(".pdf"):
            return render_template("index.html", error="Only PDF files are allowed.")

        try:
            filename = secure_filename(file.filename)
            source_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
            file.save(source_path)

            workbook = build_workbook_from_pdf(source_path)
            excel_buffer = BytesIO()
            workbook.save(excel_buffer)
            excel_buffer.seek(0)

            output_name = Path(filename).stem + "_converted.xlsx"
            return send_file(
                excel_buffer,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name=output_name,
            )
        except Exception as exc:
            return render_template("index.html", error=f"Conversion failed: {exc}")

    return render_template("index.html", error=None)


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
